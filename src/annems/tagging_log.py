"""Importer for the manual Creative Tagging Log (docs/*.xlsx).

Why this exists: until Milestone 4 is live, every ad Aiman publishes is tagged
by hand in the spreadsheet. That spreadsheet is the only record linking a Meta
ad to its concept and angle, and `angle_performance()` needs that link to tell
a winning angle from a losing one. This module loads it into `creatives`.

Discipline is enforced, not assumed: a row whose concept is not in
config/brand.yaml is rejected with a message naming the row, because a typo'd
concept silently splits one angle's data into two and teaches the system a lie.

Note on scope: the "Catatan / Learning" column is deliberately NOT imported.
Learning records belong to Agent 6, which is Phase 2 and gated on the Day-90
review (CLAUDE.md section 4). The column stays in the spreadsheet until then.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .agents.ad_writer import brand

log = logging.getLogger(__name__)

SHEET_NAME = "Tagging Log"
EXAMPLE_MARKER = "contoh"  # the seeded example row says "Contoh baris — padam..."

COLUMNS = {
    "tarikh_mula": "Tarikh Mula",
    "nama_ad": "Nama Ad (Meta)",
    "kempen": "Kempen",
    "konsep": "Konsep",
    "angle": "Angle",
    "format": "Format",
    "awareness": "Awareness",
    "status": "Status",
    "tarikh_tamat": "Tarikh Tamat",
    "keputusan": "Keputusan",
    "catatan": "Catatan / Learning",
}


@dataclass
class ImportReport:
    total_rows: int = 0
    imported: int = 0
    skipped_example: int = 0
    skipped_blank: int = 0
    rejected: list[str] = field(default_factory=list)
    unknown_concepts: set[str] = field(default_factory=set)

    def __str__(self) -> str:
        lines = [
            f"Baris dibaca      : {self.total_rows}",
            f"Diimport          : {self.imported}",
            f"Baris contoh      : {self.skipped_example}",
            f"Baris kosong      : {self.skipped_blank}",
            f"Ditolak           : {len(self.rejected)}",
        ]
        if self.unknown_concepts:
            lines.append(
                "Konsep tak dikenali: "
                + ", ".join(sorted(self.unknown_concepts))
                + "  → betulkan ejaan dalam spreadsheet, atau tambah ke config/brand.yaml"
            )
        lines += [f"  ✗ {r}" for r in self.rejected[:20]]
        if len(self.rejected) > 20:
            lines.append(f"  … dan {len(self.rejected) - 20} lagi")
        return "\n".join(lines)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def read_rows(path: str | Path) -> list[dict]:
    """Read the Tagging Log sheet into dicts keyed by our internal names."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Sheet {SHEET_NAME!r} tiada dalam {path}. Sheet yang ada: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [_cell(h) for h in rows[0]]
    index = {}
    for key, label in COLUMNS.items():
        if label in header:
            index[key] = header.index(label)
        else:
            log.warning("Column %r not found in the Tagging Log header", label)

    out = []
    for raw in rows[1:]:
        out.append({key: _cell(raw[i]) if i < len(raw) else "" for key, i in index.items()})
    return out


def validate(row: dict, row_no: int) -> str | None:
    """Return an error string, or None if the row is usable."""
    b = brand()
    if row.get("konsep") not in b["concepts"]:
        return f"Baris {row_no}: konsep {row.get('konsep')!r} tak dikenali"
    if row.get("format") and row["format"] not in b["formats"]:
        return f"Baris {row_no}: format {row.get('format')!r} tak dikenali"
    if row.get("awareness") and row["awareness"] not in b["awareness_levels"]:
        return f"Baris {row_no}: awareness {row.get('awareness')!r} tak dikenali"
    if not row.get("nama_ad"):
        return f"Baris {row_no}: 'Nama Ad (Meta)' kosong — tak boleh dipadan dengan data Meta"
    if not row.get("angle"):
        return f"Baris {row_no}: angle kosong"
    return None


def to_creative_row(row: dict) -> dict:
    return {
        "concept": row["konsep"],
        "angle": row["angle"],
        "format": row.get("format") or "Static 4:5",
        "awareness_level": row.get("awareness") or "Problem Aware",
        "meta_ad_name": row["nama_ad"],
        "created_by": "tagging_log_import",
        "qa_status": "MANUAL",
    }


def import_file(path: str | Path, dry_run: bool = False) -> ImportReport:
    """Load the spreadsheet into `creatives`. Safe to re-run."""
    from . import db

    report = ImportReport()
    rows = read_rows(path)
    report.total_rows = len(rows)

    payloads: list[dict] = []
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        if not any(row.values()):
            report.skipped_blank += 1
            continue
        if EXAMPLE_MARKER in row.get("catatan", "").lower():
            report.skipped_example += 1
            continue

        error = validate(row, i)
        if error:
            report.rejected.append(error)
            if row.get("konsep") and row["konsep"] not in brand()["concepts"]:
                report.unknown_concepts.add(row["konsep"])
            continue

        payloads.append(to_creative_row(row))

    if dry_run:
        log.info("[dry-run] %d rows would be imported", len(payloads))
        report.imported = len(payloads)
        return report

    for payload in payloads:
        try:
            db.insert_creative(payload)
            report.imported += 1
        except Exception as exc:
            report.rejected.append(f"{payload['meta_ad_name']}: {str(exc)[:120]}")
    return report
